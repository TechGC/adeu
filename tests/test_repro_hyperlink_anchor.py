import io

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine


def _doc_with_bold_hyperlink() -> io.BytesIO:
    """A paragraph whose middle is a bold, hyperlinked email — the real LLO-924 shape."""
    doc = Document()
    p = doc.add_paragraph("Email the team at ")
    rel_id = doc.part.relate_to(
        "mailto:legal@adeu.com",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "legal@adeu.com"
    r.append(t)
    hyperlink.append(r)
    p._element.append(hyperlink)
    p.add_run(" before Friday.")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


def test_anchor_resolves_across_bold_hyperlink():
    """A clean anchor that crosses a bold + hyperlinked span must resolve and apply.

    The extractor anchors on plain text (``...at legal@adeu.com before...``) but the doc renders
    the span as ``[**legal@adeu.com**](mailto:...)``. Before LLO-924 the clean matching view kept
    that markup, so the anchor was reported ANCHOR_NOT_FOUND. clean_view now drops the bracket/url
    and bold markers, so the anchor resolves and a tracked change is produced.
    """
    engine = RedlineEngine(_doc_with_bold_hyperlink(), author="QA")
    engine.apply_edits(
        [
            ModifyText(
                target_text="Email the team at legal@adeu.com before Friday.",
                new_text="Please email the team at legal@adeu.com by Friday.",
            )
        ]
    )
    xml = Document(engine.save_to_stream()).paragraphs[0]._element.xml
    assert "<w:ins" in xml, "anchor crossing the bold hyperlink did not resolve (no tracked change applied)"
