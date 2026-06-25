import io
import re

from docx import Document

from adeu.diff import trim_common_context
from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine


def _remainder(target: str, new_val: str) -> tuple[str, str]:
    p, s = trim_common_context(target, new_val)
    t_end = len(target) - s
    n_end = len(new_val) - s
    return target[p:t_end], new_val[p:n_end]


def test_trailing_punctuation_swap_trims_to_punctuation():
    """A pure trailing-punctuation swap must redline only the punctuation, not the whole word."""
    assert _remainder("jair.", "jair;") == (".", ";")
    assert _remainder("shall.", "shall,") == (".", ",")


def test_word_with_punctuation_plus_clause_preserves_word():
    """Adding a clause after a word that only changes its trailing punctuation must keep the word."""
    target = "Cause."
    new_val = "Cause; provided that Involuntary Termination shall not include a termination"
    del_text, ins_text = _remainder(target, new_val)
    assert del_text == "."
    assert ins_text.startswith(";")
    assert "Cause" not in del_text


def test_midword_change_still_redlines_whole_word():
    """Guard: a mid-word change (same character class) must still redline the whole token."""
    assert _remainder("color", "colour") == ("color", "colour")


def test_space_bounded_changes_unchanged():
    """Guard: space-delimited word swaps keep their existing prefix/suffix trimming."""
    assert trim_common_context("Context A Context", "Context B Context") == (8, 8)
    assert trim_common_context("Hello World", "Hello User") == (6, 0)


def test_word_replacement_keeps_shared_trailing_punctuation():
    """Guard: replacing the word keeps its shared trailing punctuation in the edit (suffix unchanged)."""
    assert _remainder("text.", "document.") == ("text.", "document.")


def test_token_internal_punctuation_not_split():
    """Guard: punctuation inside numbers/URLs is not a boundary, so the token stays whole."""
    assert _remainder("$10,000.00", "$15,000.00") == ("$10,000.00", "$15,000.00")
    assert _remainder("https://old-site.com/x", "https://new-site.com/x") == (
        "https://old-site.com/x",
        "https://new-site.com/x",
    )


def _apply(paragraph_text: str, target: str, new: str) -> str:
    doc = Document()
    doc.add_paragraph(paragraph_text)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    engine = RedlineEngine(stream)
    engine.apply_edits([ModifyText(target_text=target, new_text=new)])
    return Document(engine.save_to_stream()).element.xml


def test_end_to_end_punctuation_swap_does_not_restate_word():
    """End-to-end: the unchanged word must not appear in any deletion run."""
    xml = _apply(
        'Involuntary Termination shall mean a termination by the Company without Cause.',
        "Cause.",
        "Cause; provided that Involuntary Termination shall not include a termination",
    )
    dels = re.findall(r"<w:delText[^>]*>(.*?)</w:delText>", xml)
    assert not any("Cause" in d for d in dels), f"word was wholesale-deleted: {dels}"
    assert "provided that Involuntary Termination" in xml


def test_end_to_end_midword_change_still_whole_word():
    """Guard end-to-end: a mid-word change still strikes the whole word."""
    xml = _apply("The colorful banner.", "colorful", "colourful")
    dels = re.findall(r"<w:delText[^>]*>(.*?)</w:delText>", xml)
    assert any("colorful" in d for d in dels), f"mid-word change was over-trimmed: {dels}"
