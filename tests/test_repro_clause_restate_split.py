import io
import re

from docx import Document

from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine


def _apply(paragraph_text: str, target: str, new: str) -> str:
    doc = Document()
    doc.add_paragraph(paragraph_text)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    engine = RedlineEngine(stream)
    engine.apply_edits([ModifyText(target_text=target, new_text=new)])
    return Document(engine.save_to_stream()).element.xml


def test_clause_restate_below_old_gate_splits_to_minimal_edits():
    """A clause-level restate below the old 100-char gate must split into minimal sub-edits.

    Mark over-anchoring repro: deleting a ~64-char clause to add a leading word and a trailing
    qualifier used to produce a full delete+reinsert (restate-ratio ~0.8) because the surgical
    split only fired above 100 chars. The shared-block gate now splits it so the unchanged clause
    is preserved and only the genuine additions are tracked.
    """
    xml = _apply(
        "The indemnifying party shall obtain approval of any such settlement not to be unreasonably withheld.",
        "approval of any such settlement not to be unreasonably withheld.",
        "written approval of any such settlement not to be unreasonably withheld, conditioned, or delayed.",
    )
    dels = re.findall(r"<w:delText[^>]*>(.*?)</w:delText>", xml)
    assert not any("settlement not to be unreasonably" in d for d in dels), (
        f"shared clause was wholesale-deleted (over-anchoring not fixed): {dels}"
    )
    assert "written" in xml, "leading insertion missing"
    assert "conditioned, or delayed" in xml, "trailing insertion missing"


def test_genuine_rewrite_is_not_oversplit():
    """A low-overlap rewrite must stay a single clean replace, not fragment into many tiny edits."""
    xml = _apply(
        "The Company shall be solely liable for all damages.",
        "The Company shall be solely liable for all damages.",
        "Each party shall be responsible for its own costs.",
    )
    dels = re.findall(r"<w:delText[^>]*>(.*?)</w:delText>", xml)
    joined = "".join(dels)
    assert "solely liable" in joined, (
        f"low-overlap rewrite was over-split / preserved instead of cleanly replaced: {dels}"
    )


def test_scattered_rewrite_does_not_fragment_into_word_soup():
    """A reworded sentence that shares many *scattered* short runs must NOT surgical-split.

    The shared words ("do not", "internally", "need", "know", "General Counsel") push
    preserved_ratio over the gate, but no single shared block dominates — splitting on the
    diff opcodes produced word-level soup (e.g. ``c[-omm-]u[-nicate about-]{+ss+}``). The
    contiguity guard (dominant-block ratio + fanout cap) forces a single clean delete block.
    """
    target = (
        "Do not discuss the Matter externally and do not communicate about it internally "
        "except on a need-to-know basis and as directed by the General Counsel."
    )
    new = (
        "Do not discuss the Matter outside of The Suite, and do not discuss it internally "
        "except with people who need to know and only as directed by the General Counsel."
    )
    xml = _apply(target, target, new)
    dels = re.findall(r"<w:delText[^>]*>(.*?)</w:delText>", xml)
    assert len(dels) <= 2, f"scattered rewrite fragmented into {len(dels)} delete runs (word soup): {dels}"
    assert any("communicate about it internally except" in d for d in dels), (
        f"the reworded span was not deleted as one contiguous block: {dels}"
    )
