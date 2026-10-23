import io
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from adeu.models import EditComment, ModifyText
from adeu.redline.engine import RedlineEngine


def test_native_comment_creation_and_linking():
    doc = Document()
    doc.add_paragraph("The quick brown fox.")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    edit = ModifyText(target_text="quick", new_text="slow", comment="Foxes are not always quick.")

    engine = RedlineEngine(stream)
    engine.apply_edits([edit])

    result_stream = engine.save_to_stream()
    doc = Document(result_stream)

    doc_xml = doc.element.xml
    assert "w:commentRangeStart" in doc_xml
    assert "w:commentRangeEnd" in doc_xml
    assert "w:commentReference" in doc_xml

    id_match = re.search(r'<w:commentRangeStart[^>]*w:id="(\d+)"', doc_xml)
    assert id_match
    comment_id = id_match.group(1)

    comments_part = None
    for rel in doc.part.rels.values():
        if rel.reltype == RT.COMMENTS:
            comments_part = rel.target_part
            break

    assert comments_part is not None
    comments_xml = comments_part.blob.decode("utf-8")
    assert "Foxes are not always quick." in comments_xml
    assert f'w:id="{comment_id}"' in comments_xml


def test_multiple_comments_ids():
    doc = Document()
    doc.add_paragraph("First sentence.")
    doc.add_paragraph("Second sentence.")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    edit1 = ModifyText(target_text="First", new_text="First The ", comment="Comment One")
    edit2 = ModifyText(target_text="Second", new_text="Second The ", comment="Comment Two")

    engine = RedlineEngine(stream)
    engine.apply_edits([edit1, edit2])

    result_stream = engine.save_to_stream()
    doc = Document(result_stream)

    comments_part = None
    for rel in doc.part.rels.values():
        if rel.reltype == RT.COMMENTS:
            comments_part = rel.target_part
            break

    comments_xml = comments_part.blob.decode("utf-8")

    assert "Comment One" in comments_xml
    assert "Comment Two" in comments_xml

    ids = re.findall(r'<w:comment\s+[^>]*w:id="(\d+)"', comments_xml)
    assert len(set(ids)) == 2


def _apply_and_read(paragraph: str, edit: ModifyText) -> tuple[str, str]:
    doc = Document()
    doc.add_paragraph(paragraph)
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    engine = RedlineEngine(stream, author="Jane Doe")
    engine.apply_edits([edit])
    result = Document(engine.save_to_stream())

    comments_part = next(rel.target_part for rel in result.part.rels.values() if rel.reltype == RT.COMMENTS)
    return result.element.xml, comments_part.blob.decode("utf-8")


def _authors_by_text(comments_xml: str) -> dict[str, str]:
    pairs = re.findall(r'<w:comment\b[^>]*w:author="([^"]*)"[^>]*>.*?<w:t[^>]*>([^<]*)</w:t>', comments_xml, re.S)
    return {text: author for author, text in pairs}


def test_comment_list_writes_one_comment_per_entry_with_its_author():
    edit = ModifyText(
        target_text="quick",
        new_text="slow",
        comment=[
            EditComment(text="Internal: foxes vary.", author="Jane Doe (Internal)"),
            EditComment(text="Vendor: please confirm.", author="Jane Doe (Vendor)"),
        ],
    )

    _doc_xml, comments_xml = _apply_and_read("The quick brown fox.", edit)

    assert _authors_by_text(comments_xml) == {
        "Internal: foxes vary.": "Jane Doe (Internal)",
        "Vendor: please confirm.": "Jane Doe (Vendor)",
    }


def test_comment_list_entries_share_the_same_anchor():
    edit = ModifyText(
        target_text="quick",
        new_text="slow",
        comment=[EditComment(text="One"), EditComment(text="Two")],
    )

    doc_xml, _comments_xml = _apply_and_read("The quick brown fox.", edit)

    starts = re.findall(r'<w:commentRangeStart[^>]*w:id="(\d+)"', doc_xml)
    ends = re.findall(r'<w:commentRangeEnd[^>]*w:id="(\d+)"', doc_xml)
    assert len(starts) == 2
    assert ends == list(reversed(starts))
    between_starts = doc_xml.split(f'w:id="{starts[0]}"')[1].split(f'w:id="{starts[1]}"')[0]
    assert "<w:t" not in between_starts


def test_comment_list_entry_without_author_uses_the_engine_author():
    edit = ModifyText(target_text="quick", new_text="slow", comment=[EditComment(text="No author given")])

    _doc_xml, comments_xml = _apply_and_read("The quick brown fox.", edit)

    assert _authors_by_text(comments_xml) == {"No author given": "Jane Doe"}


def test_string_comment_still_writes_a_single_comment_under_the_engine_author():
    edit = ModifyText(target_text="quick", new_text="slow", comment="Plain comment")

    _doc_xml, comments_xml = _apply_and_read("The quick brown fox.", edit)

    assert _authors_by_text(comments_xml) == {"Plain comment": "Jane Doe"}


def test_surgically_split_edit_keeps_every_comment_together():
    edit = ModifyText(
        target_text="The Supplier shall deliver all of the ordered goods to the Customer within thirty days.",
        new_text="The Vendor shall deliver all of the ordered goods to the Customer within sixty days.",
        comment=[EditComment(text="Internal note", author="A"), EditComment(text="Vendor note", author="B")],
    )

    doc_xml, comments_xml = _apply_and_read(
        "The Supplier shall deliver all of the ordered goods to the Customer within thirty days.", edit
    )

    assert len(re.findall(r"<w:ins\b", doc_xml)) > 1
    assert _authors_by_text(comments_xml) == {"Internal note": "A", "Vendor note": "B"}
    starts = re.findall(r'<w:commentRangeStart[^>]*w:id="(\d+)"', doc_xml)
    assert len(starts) == 2
    between_starts = doc_xml.split(f'w:id="{starts[0]}"')[1].split(f'w:id="{starts[1]}"')[0]
    assert "<w:t" not in between_starts


def test_comment_text_joins_a_comment_list():
    edit = ModifyText(target_text="a", new_text="b", comment=[EditComment(text="One"), EditComment(text="Two")])

    assert edit.comment_text == "One\n\nTwo"


def test_paragraphs_added_after_the_target_keep_their_comment():
    sentence = "This Agreement is governed by the laws of New York."
    edit = ModifyText(
        target_text=sentence,
        new_text=(
            f"{sentence}\nConfidentiality: Firm keeps Client information confidential."
            "\nNon-Solicitation: Firm will not recruit Client employees."
        ),
        comment="Adds the missing protections.",
    )

    doc_xml, comments_xml = _apply_and_read(sentence, edit)

    assert _authors_by_text(comments_xml) == {"Adds the missing protections.": "Jane Doe"}
    start_id = re.search(r'<w:commentRangeStart[^>]*w:id="(\d+)"', doc_xml).group(1)
    anchored = doc_xml.split("<w:commentRangeStart")[1].split(f'<w:commentRangeEnd w:id="{start_id}"')[0]
    assert "Confidentiality:" in anchored
    assert "Non-Solicitation:" in anchored
