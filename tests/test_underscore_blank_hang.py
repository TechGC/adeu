import io
import re
import signal
from contextlib import contextmanager

import pytest
from docx import Document

from adeu.diff import trim_common_context
from adeu.models import ModifyText
from adeu.redline.engine import RedlineEngine

# A terminating trim_common_context returns in microseconds, so anything past this is a live lock.
_HANG_BUDGET_SECONDS = 5.0
_SIGNATURE_BLANK = "_" * 27


@contextmanager
def _fails_if_slower_than(seconds: float):
    """Turn a non-terminating call into a fast assertion failure instead of a hung test session."""

    def _raise(signum, frame):
        raise AssertionError(f"call did not return within {seconds}s — infinite loop")

    previous = signal.signal(signal.SIGALRM, _raise)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def _trim(target: str, new_val: str) -> tuple[int, int]:
    with _fails_if_slower_than(_HANG_BUDGET_SECONDS):
        return trim_common_context(target, new_val)


def test_odd_underscore_run_inside_link_terminates():
    """An odd fill-in blank inside a markdown link must not live-lock the prefix backtrack."""
    target = "email it to [_________](mailto:x@y.com) with a brief"
    _trim(target, target + " X")


def test_signature_blank_edit_terminates():
    """The reported production case: a signature-block blank shared by target and replacement."""
    _trim(
        "Name: _________ Title: Chief Executive Officer",
        "Name: _________ Title: President",
    )


def test_long_signature_blank_terminates():
    """A 27-underscore signature rule — the shape most common in real member documents."""
    _trim(f"Name: {_SIGNATURE_BLANK} Date: 2026", f"Name: {_SIGNATURE_BLANK} Date: 2027")


def test_odd_underscore_run_in_shared_suffix_terminates():
    """The suffix backtrack has the same defect: an odd blank in the shared tail must terminate."""
    _trim(f"X trailing {'_' * 9} tail", f"Y trailing {'_' * 9} tail")


@pytest.mark.parametrize("run_length", range(1, 26))
def test_underscore_runs_of_every_length_terminate(run_length: int):
    """Every blank length terminates, in both a shared prefix and a shared suffix."""
    blank = "_" * run_length
    _trim(f"Name: {blank} Title: CEO", f"Name: {blank} Title: President")
    _trim(f"intro {blank} tail same", f"intro2 {blank} tail same")


def test_balanced_italic_markers_still_backtrack():
    """Guard: a genuine lone-underscore italic pair keeps its existing balancing behaviour."""
    assert _trim("this is _italic_ text here", "this is _italic_ words here") == trim_common_context(
        "this is _italic_ text here", "this is _italic_ words here"
    )
    prefix, _suffix = _trim("a _b_ c _d_ e", "a _b_ c _d_ f")
    left = "a _b_ c _d_ e"[:prefix]
    assert left.replace("__", "").count("_") % 2 == 0


def test_even_underscore_run_unchanged():
    """Guard: even-length runs already terminated and must trim exactly as before."""
    assert _trim("Name: ____ x", "Name: ____ y") == (11, 0)


def test_dunder_token_unchanged():
    """Guard: `__dunder__` tokens stay balanced through the double-underscore branch."""
    assert _trim("call __init__ now", "call __init__ later") == (14, 0)


def test_end_to_end_signature_blank_edit_applies():
    """End-to-end: the reported docx paragraph must redline instead of hanging `process_batch`."""
    doc = Document()
    doc.add_paragraph("Name: _________ Title: Chief Executive Officer")
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    engine = RedlineEngine(stream)

    with _fails_if_slower_than(_HANG_BUDGET_SECONDS):
        result = engine.process_batch(
            [
                ModifyText(
                    target_text="Name: _________ Title: Chief Executive Officer",
                    new_text="Name: _________ Title: President",
                )
            ]
        )

    assert result["edits_applied"] == 1
    xml = Document(engine.save_to_stream()).element.xml
    assert "President" in xml
    assert re.search(r"<w:delText[^>]*>[^<]*Chief Executive Officer", xml)
